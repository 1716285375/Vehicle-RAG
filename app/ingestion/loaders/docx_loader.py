from pathlib import Path

from app.ingestion.loaders.base import BaseLoader
from app.models import Page


class DocxLoader(BaseLoader):
    async def load(self, file_path: Path) -> list[Page]:
        from docx import Document

        doc = Document(file_path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    paragraphs.append(" | ".join(cells))
        return [Page(page_number=1, text="\n\n".join(paragraphs), metadata={"source": str(file_path)})]

